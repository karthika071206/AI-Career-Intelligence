import os
import re
import csv
import io
import json
import secrets
import sqlite3
from functools import wraps
from pathlib import Path
from datetime import timedelta

from dotenv import load_dotenv
from flask import Flask, render_template, request, redirect, url_for, session, flash, g, make_response
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from ml.predict import predict_careers
from ml.predict_multifeature import predict_careers as predict_multifeature_careers
import spacy

try:
    from docx import Document
except ImportError:
    Document = None

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
# Load the English NLP model for resume analysis.
try:
    NLP = spacy.load("en_core_web_sm")
except OSError:
    NLP = None
    app.logger.warning(
        "spaCy English model is unavailable. "
        "Named entity extraction will be disabled."
    )
app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY") or secrets.token_hex(32),
    MAX_CONTENT_LENGTH=5 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
)
DATABASE = BASE_DIR / "career_intelligence.db"
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Transparent prototype skill taxonomy. Scores are heuristic suitability scores,
# not calibrated probabilities from a trained model.
CAREER_PROFILES = {
    "Software Developer": {"skills": ["Python", "Java", "C", "C++", "JavaScript", "HTML", "CSS", "SQL", "Git", "Data Structures and Algorithms", "Problem Solving"], "resources": ["Practice arrays, strings, and data structures", "Build a CRUD web application", "Learn Git and unit testing"]},
    "Data Analyst": {"skills": ["Python", "SQL", "Excel", "Power BI", "Tableau", "Pandas", "NumPy", "Data Visualization", "Statistics", "Communication"], "resources": ["Learn SQL joins and window functions", "Build a dashboard with a public dataset", "Practice statistics and data storytelling"]},
    "Data Scientist": {"skills": ["Python", "SQL", "Pandas", "NumPy", "Statistics", "Machine Learning", "Data Visualization", "Scikit-learn", "Problem Solving"], "resources": ["Study probability, statistics, and feature engineering", "Complete an end-to-end prediction project", "Learn model evaluation and experiment tracking"]},
    "AI / ML Engineer": {"skills": ["Python", "Machine Learning", "Deep Learning", "TensorFlow", "PyTorch", "NumPy", "Scikit-learn", "SQL", "Git", "Problem Solving"], "resources": ["Learn supervised and unsupervised learning", "Build and evaluate a model using scikit-learn", "Explore neural networks and model deployment"]},
    "Backend Developer": {"skills": ["Python", "Java", "SQL", "Flask", "Django", "REST API", "Git", "Data Structures and Algorithms", "Problem Solving"], "resources": ["Build REST APIs with validation", "Practice SQL schema design", "Learn authentication, testing, and deployment"]},
    "Frontend Developer": {"skills": ["HTML", "CSS", "JavaScript", "React", "Git", "Communication", "Problem Solving"], "resources": ["Create responsive pages with HTML and CSS", "Learn JavaScript and DOM fundamentals", "Build an accessible portfolio project"]},
    "Cybersecurity Analyst": {"skills": ["Networking", "Linux", "Python", "SQL", "Security", "Problem Solving", "Communication"], "resources": ["Learn networking and Linux fundamentals", "Study OWASP web security basics", "Use legal beginner labs and document findings"]},
    "Product Manager": {"skills": ["Communication", "Leadership", "Problem Solving", "Data Analysis", "SQL", "Teamwork", "Project Management"], "resources": ["Practice writing product requirements", "Learn prioritization and user research", "Analyze product metrics using a small dataset"]},
}

SKILL_PATTERNS = {
    "Python": r"\bpython\b",
    "Java": r"\bjava\b",

    # Match C++ without accidentally matching it as C.
    "C++": r"(?<![A-Za-z0-9])C\+\+(?![A-Za-z0-9])",
    "C": r"(?<![A-Za-z0-9])C(?![A-Za-z0-9+#])",

    "JavaScript": r"\bjavascript\b",
    "HTML": r"\bhtml\b",
    "CSS": r"\bcss\b",
    "SQL": r"\bsql\b",

    "Flask": r"\bflask\b",
    "Django": r"\bdjango\b",
    "Machine Learning": r"\bmachine learning\b|\bml\b",
    "Deep Learning": r"\bdeep learning\b",
    "Artificial Intelligence": r"\bartificial intelligence\b|\bAI\b",

    "Data Analysis": r"\bdata analysis\b|\bdata analytics\b",
    "Pandas": r"\bpandas\b",
    "NumPy": r"\bnumpy\b",
    "TensorFlow": r"\btensorflow\b",
    "PyTorch": r"\bpytorch\b",
    "Scikit-learn": r"\bscikit[ -]?learn\b|\bsklearn\b",

    "Power BI": r"\bpower bi\b",
    "Tableau": r"\btableau\b",
    "Git": r"\bgit\b",
    "GitHub": r"\bgithub\b",

    # Support both the abbreviation and its full form.
    "Data Structures and Algorithms":
        r"\bdata structures and algorithms\b|\bDSA\b",

    "Communication": r"\bcommunication\b",
    "Leadership": r"\bleadership\b",
    "Problem Solving": r"\bproblem[ -]solving\b",
    "Statistics": r"\bstatistics?\b",
    "Excel": r"\bexcel\b",
    "Data Visualization": r"\bdata visualization\b",

    "React": r"\breact(?:\.js)?\b",
    "REST API": r"\bREST APIs?\b|\bRESTful\b",
    "Networking": r"\bnetworking\b",
    "Linux": r"\blinux\b",
    "Security": r"\bcybersecurity\b|\bsecurity\b",
    "Teamwork": r"\bteamwork\b|\bteam player\b",
    "Project Management": r"\bproject management\b",
}


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT, full_name TEXT NOT NULL,
        email TEXT NOT NULL UNIQUE, password_hash TEXT NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    db.execute("""CREATE TABLE IF NOT EXISTS profiles (
        id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL UNIQUE,
        education TEXT DEFAULT '', branch TEXT DEFAULT '', college TEXT DEFAULT '',
        graduation_year INTEGER, career_interest TEXT DEFAULT '', technical_skills TEXT DEFAULT '',
        soft_skills TEXT DEFAULT '', updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE)""")
    # Keep resume text server-side rather than inside Flask's client-side cookie.
    db.execute("""CREATE TABLE IF NOT EXISTS resume_context (
        user_id INTEGER PRIMARY KEY,
        filename TEXT NOT NULL DEFAULT '',
        resume_text TEXT NOT NULL DEFAULT '',
        education_json TEXT NOT NULL DEFAULT '[]',
        job_titles_json TEXT NOT NULL DEFAULT '[]',
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE)""")
    db.commit()


with app.app_context():
    init_db()


@app.before_request
def csrf_protect():
    if request.method == "POST":
        submitted = request.form.get("csrf_token", "")
        expected = session.get("csrf_token", "")
        if not expected or not secrets.compare_digest(submitted, expected):
            return "Security token missing or expired. Refresh the page and try again.", 400


@app.context_processor
def inject_template_values():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"]}


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if "user_id" not in session:
            flash("Please sign in to continue.", "error")
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped_view


def get_user_and_profile():
    db = get_db()
    user = db.execute("SELECT id, full_name, email FROM users WHERE id=?", (session["user_id"],)).fetchone()
    profile = db.execute("SELECT * FROM profiles WHERE user_id=?", (session["user_id"],)).fetchone()
    return user, profile


def split_skills(value):
    return {part.strip().lower() for part in re.split(r"[,;\n]", value or "") if part.strip()}
def prepare_ml_skills(skills):
    """Convert a list of skills into text for the trained ML model."""
    cleaned = []
    seen = set()

    for skill in skills:
        skill = str(skill).strip()

        if skill and skill.lower() not in seen:
            cleaned.append(skill)
            seen.add(skill.lower())

    return ", ".join(cleaned)

def canonical_skills(text):
    return [skill for skill, pattern in SKILL_PATTERNS.items() if re.search(pattern, text or "", re.IGNORECASE)]
def extract_resume_information(text):
    """
    Extract categorized skills, education, job-title keywords,
    and named entities from resume text.
    """

    text = text or ""

    # 1. Extract skills using the existing skill patterns.
    detected_skills = canonical_skills(text)

    technical_skill_names = {
        "Python",
        "Java",
        "C",
        "C++",
        "JavaScript",
        "HTML",
        "CSS",
        "SQL",
        "Flask",
        "Django",
        "Machine Learning",
        "Deep Learning",
        "Artificial Intelligence",
        "Data Analysis",
        "Pandas",
        "NumPy",
        "TensorFlow",
        "PyTorch",
        "Scikit-learn",
        "Power BI",
        "Tableau",
        "Git",
        "GitHub",
        "Data Structures and Algorithms",
        "Statistics",
        "Excel",
        "Data Visualization",
        "React",
        "REST API",
        "Networking",
        "Linux",
        "Security",
    }

    soft_skill_names = {
        "Communication",
        "Leadership",
        "Problem Solving",
        "Teamwork",
        "Project Management",
    }

    technical_skills = [
        skill for skill in detected_skills
        if skill in technical_skill_names
    ]

    soft_skills = [
        skill for skill in detected_skills
        if skill in soft_skill_names
    ]

    # 2. Extract certification/course names using keyword matching.
    # This detects mentions; it does not verify that a credential was earned.
    certification_patterns = [
        r"Programming Concepts in Python",
        r"Getting Started with Object-Oriented Programming in Python 3",
        r"Modernizing Python Applications: Migrating to Python 3\.x",
        r"Real World Projects in Python 3\.x",
        r"Fast Numerical Computing with Python",
        r"Python for Data Science",
        r"Deep Learning for Developers",
        r"Generative Models for Developers",
        r"Infosys Springboard",
        r"upGrad",
        r"Juvi",
        r"AWS Certified[^\n.;|]{0,100}",
        r"Microsoft Certified[^\n.;|]{0,100}",
        r"Azure Fundamentals",
        r"Google Cloud Certified[^\n.;|]{0,100}",
        r"Google Data Analytics",
        r"TensorFlow Developer Certificate",
        r"Certified Kubernetes Administrator",
        r"Oracle Certified[^\n.;|]{0,100}",
        r"Cisco Certified[^\n.;|]{0,100}",
        r"CompTIA Security\+",
        r"Professional Scrum Master",
        r"Project Management Professional",
        r"\bPMP\b",
        r"\bJava Programming\b",
        r"\bMachine Learning Certification\b",
        r"\bData Science Certification\b",
        r"Java Programming Course(?:\s+Guvi)?",
        r"\bGuvi\b",
        r"Infosys Pragathi Cohort\s*\d*",
        r"Young Professional Certification",
        r"AI certificate course from upGrad",
        r"\bInfosys Springboard\b",
    ]
    certifications = []
    seen_certifications = set()

    # First, detect known certification/course names wherever they appear.
    for pattern in certification_patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            certification = re.sub(r"\s+", " ", match.group(0)).strip(" -:|•")
            key = certification.casefold()
            if certification and key not in seen_certifications:
                certifications.append(certification)
                seen_certifications.add(key)

    # Collect certificate names without including long descriptions.
    lines = [
        re.sub(r"\s+", " ", line).strip()
        for line in text.splitlines()
    ]

    in_certification_section = False

    section_heading = re.compile(
        r"^(certifications?|certificates?|professional certifications?|"
        r"courses?|training|licenses and certifications)\s*:?\s*$",
        re.IGNORECASE,
    )

    other_heading = re.compile(
        r"^(education|academic background|experience|work experience|"
        r"employment|projects|technical skills|skills|soft skills|"
        r"achievements(?:\s*(?:&|and)\s*activities)?|awards|interests|"
        r"languages|references|summary|objective|profile|contact|"
        r"publications|internships|positions of responsibility)\s*:?\s*$",
        re.IGNORECASE,
    )

    description_words = re.compile(
        r"\b(finished|completed|developed|covering|contributing|"
        r"strengthened|competed|understanding|career readiness|"
        r"under time constraints|gained experience|learned how to|"
        r"responsible for|fundamentals|problem-solving using)\b",
        re.IGNORECASE,
    )

    for line in lines:
        cleaned = re.sub(r"^[\s•●▪\-*|]+", "", line).strip()
        cleaned = re.sub(r"^[0-9]+[.)]\s*", "", cleaned).strip()

        if not cleaned:
            continue

        if section_heading.fullmatch(cleaned):
            in_certification_section = True
            continue

        if other_heading.fullmatch(cleaned):
            in_certification_section = False
            continue

        if not in_certification_section:
            continue

        # Ignore descriptions, sentences, and excessively long entries.
        if len(cleaned) > 100:
            continue

        if description_words.search(cleaned):
            continue

        if cleaned.count(".") > 1:
            continue

        candidate = cleaned.strip(" -:|•")

        if not candidate:
            continue

        key = candidate.casefold()

        if key not in seen_certifications:
            certifications.append(candidate)
            seen_certifications.add(key)

    # 2. Identify common education terms.
    education_patterns = [
        r"\bB\.?\s?Tech\b",
        r"\bB\.?\s?E\.?\b",
        r"\bM\.?\s?Tech\b",
        r"\bM\.?\s?E\.?\b",
        r"\bB\.?\s?Sc\b",
        r"\bM\.?\s?Sc\b",
        r"\bMBA\b",
        r"\bMCA\b",
        r"\bBCA\b",
        r"\bBachelor(?:'s)?\b",
        r"\bMaster(?:'s)?\b",
        r"\bPh\.?\s?D\.?\b",
        r"\bIntermediate\b",
        r"\bHigher Secondary\b",
        r"\bSSC\b",
        r"\bClass\s+(?:10|12|X|XII)\b",
    ]

    education = []

    for pattern in education_patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            value = match.group(0).strip()

            if value.lower() not in {
                item.lower() for item in education
            }:
                education.append(value)

    # 3. Find common job-title keywords.
    job_title_patterns = [
        r"\bSoftware Engineer\b",
        r"\bSoftware Developer\b",
        r"\bPython Developer\b",
        r"\bJava Developer\b",
        r"\bWeb Developer\b",
        r"\bFrontend Developer\b",
        r"\bBackend Developer\b",
        r"\bFull Stack Developer\b",
        r"\bData Analyst\b",
        r"\bData Scientist\b",
        r"\bMachine Learning Engineer\b",
        r"\bAI Engineer\b",
        r"\bAI/ML Engineer\b",
        r"\bIntern\b",
        r"\bSoftware Engineering Intern\b",
        r"\bData Analyst Intern\b",
        r"\bbusiness analyst\b",
    ]

    job_titles = []

    for pattern in job_title_patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            value = match.group(0).strip()

            if value.lower() not in {
                item.lower() for item in job_titles
            }:
                job_titles.append(value)

    # 4. Recognize named entities using spaCy.
    entities = []

    if NLP is not None:
        doc = NLP(text[:100000])

        for ent in doc.ents:
            if ent.label_ in {
                "PERSON",
                "ORG",
                "DATE",
                "GPE",
            }:
                entities.append({
                    "text": ent.text,
                    "label": ent.label_,
                })

    return {
        "technical_skills": technical_skills,
        "soft_skills": soft_skills,
        "education": education,
        "job_titles": job_titles,
        "certifications": certifications,
        "entities": entities,
    }

def build_recommendations(
    skills,
    interests="",
    education=None,
    job_titles=None,
):
    """
    Rank careers using technical skills, education,
    previous job titles, and career interests.

    The score is a heuristic suitability score, not a
    calibrated probability of employment or career success.
    """

    education = education or []
    job_titles = job_titles or []

    # Normalize candidate information.
    user_skills = {
        str(skill).strip().casefold()
        for skill in skills
        if str(skill).strip()
    }

    interest_text = str(interests or "").casefold()

    education_text = " ".join(
        str(item) for item in education
    ).casefold()

    job_title_text = " ".join(
        str(item) for item in job_titles
    ).casefold()

    results = []

    # Career-related keywords provide a small, transparent
    # contextual signal. They are not learned ML features.
    career_keywords = {
        "Software Developer": [
            "software developer",
            "software engineer",
            "programming",
            "computer science",
        ],
        "Data Analyst": [
            "data analyst",
            "business analyst",
            "data analysis",
            "analytics",
        ],
        "Data Scientist": [
            "data scientist",
            "data science",
            "statistics",
            "machine learning",
        ],
        "AI / ML Engineer": [
            "ai engineer",
            "machine learning engineer",
            "artificial intelligence",
            "deep learning",
        ],
        "Backend Developer": [
            "backend developer",
            "back-end developer",
            "api developer",
            "server-side",
        ],
        "Frontend Developer": [
            "frontend developer",
            "front-end developer",
            "web designer",
            "ui developer",
        ],
        "Cybersecurity Analyst": [
            "cybersecurity analyst",
            "security analyst",
            "information security",
            "network security",
        ],
        "Product Manager": [
            "product manager",
            "product management",
            "product owner",
        ],
    }

    technical_education_terms = [
        "computer science",
        "information technology",
        "artificial intelligence",
        "machine learning",
        "data science",
        "software engineering",
        "computer engineering",
    ]

    technical_careers = {
        "Software Developer",
        "Data Analyst",
        "Data Scientist",
        "AI / ML Engineer",
        "Backend Developer",
        "Frontend Developer",
        "Cybersecurity Analyst",
    }

    for career, profile in CAREER_PROFILES.items():
        required = profile["skills"]

        # 1. Match technical and soft skills.
        matched = [
            skill
            for skill in required
            if skill.casefold() in user_skills
        ]

        missing = [
            skill
            for skill in required
            if skill.casefold() not in user_skills
        ]

        skill_score = (
            len(matched) / max(1, len(required))
        ) * 70

        # 2. Consider stated career interests.
        keywords = career_keywords.get(
            career,
            [career.casefold()],
        )

        interest_match = any(
            keyword in interest_text
            for keyword in keywords
        )

        interest_score = 10 if interest_match else 0

        # 3. Consider previous job-title keywords.
        experience_match = any(
            keyword in job_title_text
            for keyword in keywords
        )

        experience_score = 10 if experience_match else 0

        # 4. Consider education conservatively.
        education_match = (
            any(
                term in education_text
                for term in technical_education_terms
            )
            and career in technical_careers
        )

        education_score = 10 if education_match else 0

        # Final heuristic suitability score.
        score = round(
            min(
                100,
                skill_score
                + interest_score
                + experience_score
                + education_score,
            ),
            2,
        )

        results.append({
            "career": career,
            "score": score,
            "matched": matched,
            "missing": missing,
            "required_count": len(required),
            "resources": profile["resources"],
            "education_match": education_match,
            "experience_match": experience_match,
            "interest_match": interest_match,
        })

    return sorted(
        results,
        key=lambda item: (
            item["score"],
            len(item["matched"]),
        ),
        reverse=True,
    )


@app.errorhandler(413)
def too_large(error):
    flash("The file is too large. Please upload a file smaller than 5 MB.", "error")
    return redirect(url_for("resume_analyzer"))


@app.route("/")
def index():
    return redirect(url_for("dashboard" if session.get("user_id") else "login"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        if not full_name or not email or not password or not confirm:
            flash("Please complete every field.", "error")
        elif len(full_name) > 100 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            flash("Enter a valid name and email address.", "error")
        elif len(password) < 10 or not re.search(r"[A-Za-z]", password) or not re.search(r"\d", password):
            flash("Use at least 10 characters, including letters and numbers.", "error")
        elif password != confirm:
            flash("The passwords do not match.", "error")
        else:
            db = get_db()
            try:
                cursor = db.execute("INSERT INTO users (full_name,email,password_hash) VALUES (?,?,?)",
                                    (full_name, email, generate_password_hash(password)))
                db.execute("INSERT OR IGNORE INTO profiles (user_id) VALUES (?)", (cursor.lastrowid,))
                db.commit()
                session.clear()
                session["user_id"] = cursor.lastrowid
                session["csrf_token"] = secrets.token_urlsafe(32)
                session.permanent = True
                flash("Your account is ready. Welcome to CareerCast!", "success")
                return redirect(url_for("dashboard"))
            except sqlite3.IntegrityError:
                db.rollback()
                flash("An account with this email already exists. Please sign in.", "error")
    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = get_db().execute("SELECT id,password_hash FROM users WHERE email=?", (email,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["csrf_token"] = secrets.token_urlsafe(32)
            session.permanent = True
            flash("Signed in successfully.", "success")
            return redirect(url_for("dashboard"))
        flash("Email or password is incorrect.", "error")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    flash("You have been signed out.", "success")
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    user, profile = get_user_and_profile()
    if user is None:
        session.clear()
        return redirect(url_for("login"))
    skills = split_skills(profile["technical_skills"] if profile else "")
    fields = ["education", "branch", "college", "graduation_year", "career_interest", "technical_skills", "soft_skills"]
    completion = round(sum(bool(profile and profile[f]) for f in fields) / len(fields) * 100)
    recommendations = build_recommendations(skills, profile["career_interest"] if profile else "")[:3]
    return render_template("dashboard.html", full_name=user["full_name"], skills_count=len(skills),
                           completion=completion, recommendations=recommendations)


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    db = get_db()
    user, current = get_user_and_profile()
    if request.method == "POST":
        values = {key: request.form.get(key, "").strip() for key in
                  ["full_name", "education", "branch", "college", "graduation_year", "career_interest", "technical_skills", "soft_skills"]}
        if not all(values.values()):
            flash("Please complete all profile fields.", "error")
        elif len(values["full_name"]) > 100:
            flash("Name must be 100 characters or fewer.", "error")
        else:
            try:
                year = int(values["graduation_year"])
                if not 2000 <= year <= 2100: raise ValueError
            except ValueError:
                flash("Enter a graduation year between 2000 and 2100.", "error")
            else:
                try:
                    db.execute("UPDATE users SET full_name=? WHERE id=?", (values["full_name"], session["user_id"]))
                    db.execute("""INSERT INTO profiles (user_id,education,branch,college,graduation_year,career_interest,technical_skills,soft_skills,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(user_id) DO UPDATE SET
                    education=excluded.education, branch=excluded.branch, college=excluded.college,
                    graduation_year=excluded.graduation_year, career_interest=excluded.career_interest,
                    technical_skills=excluded.technical_skills, soft_skills=excluded.soft_skills, updated_at=CURRENT_TIMESTAMP""",
                    (session["user_id"], values["education"], values["branch"], values["college"], year,
                     values["career_interest"], values["technical_skills"], values["soft_skills"]))
                    db.commit()
                    flash("Profile saved successfully.", "success")
                    return redirect(url_for("profile"))
                except sqlite3.Error:
                    db.rollback(); app.logger.exception("Profile save failed")
                    flash("Could not save your profile. Check the VS Code terminal for details.", "error")
        user, current = get_user_and_profile()
    return render_template("profile.html", user=user, profile=current)


@app.route("/resume-analyzer", methods=["GET", "POST"])
@login_required
def resume_analyzer():
    result = None
    if request.method == "POST":
        uploaded = request.files.get("resume")
        if not uploaded or not uploaded.filename:
            flash("Choose a PDF or DOCX resume first.", "error")
            return render_template("resume.html", result=None)
        filename = secure_filename(uploaded.filename)
        ext = Path(filename).suffix.lower()
        if not filename or ext not in ALLOWED_EXTENSIONS:
            flash("Only PDF and DOCX files are supported.", "error")
            return render_template("resume.html", result=None)
        try:
            if ext == ".pdf":
                reader = PdfReader(uploaded.stream)
                if reader.is_encrypted: raise ValueError("This PDF is password-protected. Upload an unlocked PDF.")
                if not reader.pages: raise ValueError("The PDF has no pages.")
                if len(reader.pages) > 20: raise ValueError("Please upload a resume with 20 pages or fewer.")
                text_content = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
            else:
                if Document is None: raise ValueError("DOCX support is not installed. Run pip install python-docx.")
                document = Document(uploaded.stream)
                text_content = "\n".join(p.text for p in document.paragraphs).strip()
                for table in document.tables:
                    text_content += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
            if not text_content:
                raise ValueError("No selectable text was found. Scanned PDFs need OCR, which is not included.")
            # Extract categorized information from the resume.
            extracted_info = extract_resume_information(text_content)
            print("DEBUG RESUME TEXT:", repr(text_content[:3000]))
            print("DEBUG CERTIFICATIONS:", extracted_info.get("certifications"))

            # Preserve the existing skills list for career recommendations.
            found = list(dict.fromkeys(
                extracted_info["technical_skills"]
                + extracted_info["soft_skills"]
                + [
                    skill for skill in canonical_skills(text_content)
                    if skill not in extracted_info["technical_skills"]
                    and skill not in extracted_info["soft_skills"]
                ]
            ))

            # Store compact extracted fields in the session.
            # Keep the full resume text in SQLite, not in the browser cookie.
            session["resume_skills"] = found
            session["resume_technical_skills"] = extracted_info["technical_skills"]
            session["resume_soft_skills"] = extracted_info["soft_skills"]
            session["resume_education"] = extracted_info["education"]
            session["resume_job_titles"] = extracted_info["job_titles"]
            session["resume_certifications"] = extracted_info.get("certifications", [])
            session["resume_filename"] = filename

            db = get_db()
            db.execute(
                """INSERT INTO resume_context
                   (user_id, filename, resume_text, education_json, job_titles_json, updated_at)
                   VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                   ON CONFLICT(user_id) DO UPDATE SET
                       filename=excluded.filename,
                       resume_text=excluded.resume_text,
                       education_json=excluded.education_json,
                       job_titles_json=excluded.job_titles_json,
                       updated_at=CURRENT_TIMESTAMP""",
                (
                    session["user_id"],
                    filename,
                    text_content,
                    json.dumps(extracted_info["education"], ensure_ascii=False),
                    json.dumps(extracted_info["job_titles"], ensure_ascii=False),
                ),
            )
            db.commit()

            # Prepare the results for the Resume Analyzer page.
            result = {
                "filename": filename,
                "text_length": len(text_content),
                "skills": found,
                "technical_skills": extracted_info["technical_skills"],
                "soft_skills": extracted_info["soft_skills"],
                "education": extracted_info["education"],
                "job_titles": extracted_info["job_titles"],
                "certifications": extracted_info.get("certifications", []),
                "entities": extracted_info["entities"],
                "text": text_content[:12000],
            }
            print("DEBUG: Certifications passed to template:", result["certifications"])
            flash("Resume parsed. Review the recognized skills before using them for recommendations.", "success")
        except ValueError as exc:
            flash(str(exc), "error")
        except Exception:
            app.logger.exception("Resume parsing failed")
            flash("Could not read this file. Try a standard, text-based PDF or DOCX resume.", "error")
    return render_template("resume.html", result=result)
@app.route("/career-recommendations")
@login_required
def career_recommendations():
    print("\n[DEBUG] career_recommendations() was called")

    print("\n" + "=" * 65)
    print("CAREERCAST - CAREER RECOMMENDATION SYSTEM")
    print("=" * 65)

    _, profile = get_user_and_profile()

    # 1. Collect skills from the user's profile.
    profile_skills = []

    if profile:
        profile_skills = [
            skill.strip()
            for skill in re.split(
                r"[,;\n]",
                profile["technical_skills"] or ""
            )
            if skill.strip()
        ]

    # 2. Collect skills extracted from the uploaded resume.
    resume_skills = session.get("resume_skills", [])

    # 3. Retrieve saved resume education and job titles.
    resume_context = get_db().execute(
        """
        SELECT filename, education_json, job_titles_json
        FROM resume_context
        WHERE user_id = ?
        """,
        (session["user_id"],),
    ).fetchone()

    resume_education = []
    resume_job_titles = []

    if resume_context:
        try:
            resume_education = json.loads(
                resume_context["education_json"] or "[]"
            )
        except (json.JSONDecodeError, TypeError):
            resume_education = []

        try:
            resume_job_titles = json.loads(
                resume_context["job_titles_json"] or "[]"
            )
        except (json.JSONDecodeError, TypeError):
            resume_job_titles = []

    # 4. Combine profile and resume skills without duplicates.
    all_skills = []
    seen_skills = set()

    for skill in profile_skills + resume_skills:
        cleaned = str(skill).strip()

        if cleaned and cleaned.casefold() not in seen_skills:
            all_skills.append(cleaned)
            seen_skills.add(cleaned.casefold())

    interest = profile["career_interest"] if profile else ""
    recommendations = []
    error_message = None

    ml_status = (
        "Trained ML model unavailable; using profile-based recommendations."
    )

    print("\nSkills collected:")
    print(", ".join(all_skills) if all_skills else "No skills found.")

    print("\nCareer interest:", interest or "Not specified")

    # 5. Generate recommendations.
    if not all_skills:
        error_message = (
            "Please add technical skills to your profile "
            "or upload a resume before requesting recommendations."
        )

        print("\nERROR:", error_message)

    else:
        # Generate the existing profile-based suitability scores.
        heuristic_results = build_recommendations(
            all_skills,
            interests=interest,
            education=resume_education,
            job_titles=resume_job_titles,
        )

        ml_rank_by_career = {}

        # 6. Use the trained career prediction model.
        try:
            skills_text = prepare_ml_skills(all_skills)

            ml_results = predict_multifeature_careers(
                skills=skills_text,
                job_description="",
                certifications=", ".join(
                    session.get("resume_certifications", [])
                ),
                top_k=118,
            )

            print("\n" + "-" * 65)
            print("TRAINED MODEL PREDICTIONS")
            print("-" * 65)

            if not ml_results:
                print("The trained model returned no predictions.")

            for rank, result in enumerate(ml_results[:10], start=1):
                print(
                    f"{rank}. {result.get('career', 'Unknown')} "
                    f"- {float(result.get('score', 0)):.2f}%"
                )

            # Map model career labels to CareerCast's existing categories.
            career_aliases = {
                "Software Developer": (
                    "software developer",
                    "software engineer",
                    "software development",
                    "programmer",
                ),
                "Data Analyst": (
                    "data analyst",
                    "business analyst",
                    "data analytics",
                    "bi analyst",
                ),
                "Data Scientist": (
                    "data scientist",
                    "data science",
                ),
                "AI / ML Engineer": (
                    "machine learning engineer",
                    "ml engineer",
                    "ai engineer",
                    "artificial intelligence engineer",
                    "deep learning engineer",
                    "ai researcher",
                    "mlops engineer",
                ),
                "Backend Developer": (
                    "backend developer",
                    "back-end developer",
                    "backend engineer",
                    "api developer",
                ),
                "Frontend Developer": (
                    "frontend developer",
                    "front-end developer",
                    "frontend engineer",
                ),
                "Cybersecurity Analyst": (
                    "cybersecurity analyst",
                    "security analyst",
                    "information security analyst",
                ),
                "Product Manager": (
                    "product manager",
                    "product owner",
                ),
            }

            for rank, result in enumerate(ml_results, start=1):
                model_career = str(
                    result.get("career", "")
                ).casefold().strip()

                if not model_career:
                    continue

                for career, aliases in career_aliases.items():
                    if any(
                        alias in model_career
                        for alias in aliases
                    ):
                        previous_rank = ml_rank_by_career.get(career)

                        if (
                            previous_rank is None
                            or rank < previous_rank
                        ):
                            ml_rank_by_career[career] = rank

            ml_status = (
                "Hybrid ranking: profile suitability plus trained-model ranking."
            )

            print("\nModel ranks mapped to CareerCast:")
            print(
                ml_rank_by_career
                if ml_rank_by_career
                else "No model predictions matched the supported career categories."
            )

        except Exception:
            app.logger.exception(
                "Trained model prediction failed."
            )

            print(
                "\nWARNING: ML prediction failed. "
                "Using profile-based recommendations."
            )

        # 7. Combine suitability scores with model ranking.
        for item in heuristic_results:
            career = item["career"]
            ml_rank = ml_rank_by_career.get(career)

            ml_bonus = (
                max(0, 10 - 2 * (ml_rank - 1))
                if ml_rank is not None
                else 0
            )

            combined = dict(item)
            combined["ml_rank"] = ml_rank
            combined["ml_support"] = ml_rank is not None

            combined["score"] = round(
                min(
                    100.0,
                    float(item["score"]) + ml_bonus,
                ),
                2,
            )

            combined["score_type"] = "Hybrid suitability score"

            recommendations.append(combined)

        # 8. Sort and keep the top three recommendations.
        recommendations.sort(
            key=lambda item: (
                item["score"],
                item["ml_support"],
                len(item["matched"]),
            ),
            reverse=True,
        )

        recommendations = recommendations[:3]

        # 9. Print final recommendations in the terminal.
        print("\n" + "=" * 65)
        print("FINAL TOP 3 CAREER RECOMMENDATIONS")
        print("=" * 65)

        for rank, item in enumerate(recommendations, start=1):
            print(f"\nMATCH #{rank}: {item['career']}")
            print(f"Suitability score: {item['score']:.2f}%")
            print(
                "Matched skills: "
                + (
                    ", ".join(item["matched"])
                    if item["matched"]
                    else "None"
                )
            )
            print(
                "Skills to develop: "
                + (
                    ", ".join(item["missing"])
                    if item["missing"]
                    else "No additional skills identified."
                )
            )
            print(
                "Supported by trained model:",
                "Yes" if item["ml_support"] else "No",
            )

    print("\nRecommendation status:", ml_status)
    print("=" * 65 + "\n")

    # 10. Display recommendations on the existing website.
    return render_template(
        "recommendations.html",
        recommendations=recommendations,
        skills=all_skills,
        interest=interest,
        resume_filename=(
            session.get("resume_filename")
            or (
                resume_context["filename"]
                if resume_context
                else None
            )
        ),
        error_message=error_message,
        ml_status=ml_status,
    )


@app.route("/skill-gap", methods=["GET", "POST"])
@login_required
def skill_gap():
    _, profile = get_user_and_profile()
    profile_skills = [s.strip() for s in re.split(r"[,;\n]", profile["technical_skills"] or "") if s.strip()] if profile else []
    all_skills = list(dict.fromkeys(profile_skills + session.get("resume_skills", [])))
    selected = request.form.get("career", "") if request.method == "POST" else request.args.get("career", "")
    if selected not in CAREER_PROFILES:
        selected = "AI / ML Engineer" if "ai" in (profile["career_interest"] or "").lower() else "Software Developer"
    required = CAREER_PROFILES[selected]["skills"]
    current_lower = {s.lower() for s in all_skills}
    matched = [s for s in required if s.lower() in current_lower]
    missing = [s for s in required if s.lower() not in current_lower]
    alignment = round(len(matched) / len(required) * 100) if required else 0
    return render_template("skill_gap.html", careers=list(CAREER_PROFILES), selected=selected, matched=matched,
                           missing=missing, alignment=alignment, skills=all_skills,
                           resources=CAREER_PROFILES[selected]["resources"])


@app.route("/export/recommendations.csv")
@login_required
def export_recommendations():
    _, profile = get_user_and_profile()
    profile_skills = [s.strip() for s in re.split(r"[,;\n]", profile["technical_skills"] or "") if s.strip()] if profile else []
    skills = list(dict.fromkeys(profile_skills + session.get("resume_skills", [])))
    recs = build_recommendations(skills, profile["career_interest"] if profile else "")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Career path", "Heuristic suitability score (%)", "Matched skills", "Skills to develop"])
    for rec in recs:
        writer.writerow([rec["career"], rec["score"], ", ".join(rec["matched"]), ", ".join(rec["missing"])])
    response = make_response("\ufeff" + output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = "attachment; filename=careercast_recommendations.csv"
    return response


@app.errorhandler(404)
def not_found(error):
    return render_template("error.html", code=404, message="That page could not be found."), 404


@app.errorhandler(500)
def server_error(error):
    app.logger.exception("Internal server error")
    return render_template("error.html", code=500, message="Something went wrong. Please try again."), 500


if __name__ == "__main__":
    # Debug is intentionally off by default. Enable FLASK_DEBUG=1 only for local development.
    app.run(host="127.0.0.1", port=int(os.getenv("PORT", "5000")), debug=os.getenv("FLASK_DEBUG") == "1")
